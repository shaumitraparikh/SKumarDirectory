import docx
import json

doc_path = "docs 2025/GSC - SK Catlog Photo.docx"
doc = docx.Document(doc_path)

print("Paragraphs in Photo docx:")
for i, p in enumerate(doc.paragraphs[:30]):
    if p.text.strip():
        print(f"P{i}: {p.text.strip()}")

for i, table in enumerate(doc.tables[:2]):
    for row in table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        print(row_data)

